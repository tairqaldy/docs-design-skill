"""tairdoc — typeset-looking documents (DOCX → PDF via LibreOffice) in a classical, academic style.

Three presets, all monochrome serif typography:
  article  LaTeX `article` look: centred title block, bold numbered headings, Libertine + CM math.
  journal  Old-school journal paper: title between hairlines, small-caps author, abstract,
           keywords, letter-spaced section heads, running footer.
  essay    Tufte handout: italic left-aligned title, Palatino, wide right margin with
           side notes (^[note] in text) and margin captions.

Minimal use (full guide: ../SKILL.md, design rules: ../references/design-guide.md):

    import sys; sys.path.insert(0, r"<skill folder>/scripts")
    from tairdoc import Doc, Diagram, render_terminal
    d = Doc("Practice 6", author="Jane Doe", affiliation="AI Fundamentals",
            date="25 October 2026", theme="article")
    d.section("Bayes' theorem")
    d.para("Of the 38 BLOCKED reports, 24 are real.^[The other 14 come from clear paths.]")
    d.equation(r"P(O\\mid B)=\\frac{0.24}{0.38}")
    d.save("practice6.docx")          # writes .docx and .pdf

Dependencies: python-docx, matplotlib, Pillow; LibreOffice for PDF; PyMuPDF for preview/contents.
"""
from __future__ import annotations

import copy
import hashlib
import re
import shutil
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Pt, RGBColor

# --------------------------------------------------------------------------- fonts
# Each preset names a first-choice face plus equivalents; the first installed one is used.
# (LibreOffice bundles Linux Libertine G, so the article/journal presets work wherever LO does.)
FALLBACKS = {
    "Linux Libertine G": ["Linux Libertine G", "Linux Libertine O", "Libertinus Serif", "Linux Libertine"],
    "Palatino Linotype": ["Palatino Linotype", "Palatino", "TeX Gyre Pagella", "P052", "URW Palladio L",
                          "Book Antiqua"],
    "Consolas": ["Consolas", "Menlo", "SF Mono", "Cascadia Mono", "DejaVu Sans Mono", "Liberation Mono"],
}


def _installed():
    try:
        from matplotlib import font_manager
        return {f.name for f in font_manager.fontManager.ttflist}
    except Exception:
        return set()


def pick_font(name):
    """Return `name` or its first installed equivalent (see FALLBACKS)."""
    have = _installed()
    for cand in FALLBACKS.get(name, [name]):
        if cand in have:
            return cand
    return name  # not visible to matplotlib; LibreOffice may still have it (e.g. its bundled fonts)


MONO = pick_font("Consolas")
INK_HEX = "111111"
INK = RGBColor(0x11, 0x11, 0x11)
GREY_HEX = "555555"
RULE_HEX = "1A1A1A"
CODE_BAR_HEX = "A0A0A0"

# --------------------------------------------------------------------------- themes
# Font names may carry OpenType/Graphite features after a colon; LibreOffice honours them:
#   onum = old-style figures, lnum/tnum = lining/tabular figures, smcp = true small caps.
THEMES = {
    "article": dict(
        serif="Linux Libertine G", prose="onum", figs="lnum&tnum", math="cm",
        body=11, leading=1.17, indent=0.6,
        margins=(3.1, 3.1, 2.7, 2.6),                 # left, right, top, bottom (cm)
        note_col=0, note_gap=0,
        title=dict(align="center", size=18, italic=False),
        h1=dict(size=13.5, bold=True, italic=False, upper=False, track=0, before=18, after=7, fmt="{n}\u2003{t}"),
        h2=dict(size=11.5, bold=True, italic=False, upper=False, track=0, before=12, after=4, fmt="{n}\u2003{t}"),
        h3=dict(size=11, bold=True, italic=True, upper=False, track=0, before=9, after=2, fmt="{t}"),
        caption=dict(label="{K} {n}:", smcp=False, italic=False, size=10),
        table_head="bold", footer="centre", bullet="\u2022",
    ),
    "journal": dict(
        serif="Linux Libertine G", prose="", figs="lnum&tnum", math="cm",
        body=10.5, leading=1.14, indent=0.5,
        margins=(2.4, 2.4, 2.4, 2.7),
        note_col=0, note_gap=0,
        title=dict(align="center", size=21, italic=False),
        h1=dict(size=9.8, bold=True, italic=False, upper=True, track=0.9, before=16, after=7, fmt="{n}.\u2002{t}"),
        h2=dict(size=10.5, bold=True, italic=True, upper=False, track=0, before=11, after=3, fmt="{n}\u2002{t}"),
        h3=dict(size=10.5, bold=False, italic=True, upper=False, track=0, before=8, after=2, fmt="{t}"),
        caption=dict(label="{K} {n}.", smcp=True, italic=False, size=9.5),
        table_head="smcp", footer="running", bullet="\u2022",
    ),
    "essay": dict(
        serif="Palatino Linotype", prose="onum", figs="lnum&tnum", math="cm",
        body=11, leading=1.3, indent=0.55,
        margins=(2.5, 7.7, 2.7, 2.6),
        note_col=5.4, note_gap=0.9,
        title=dict(align="left", size=22, italic=True),
        h1=dict(size=15, bold=False, italic=True, upper=False, track=0, before=20, after=7, fmt="{t}"),
        h2=dict(size=12, bold=False, italic=True, upper=False, track=0, before=13, after=4, fmt="{t}"),
        h3=dict(size=11, bold=True, italic=False, upper=False, track=0, before=9, after=2, fmt="{t}"),
        caption=dict(label="{K} {n}.", smcp=False, italic=True, size=9),
        table_head="italic", footer="right", bullet="\u2013",
    ),
}

_TMP = Path(tempfile.gettempdir())
_ICON_DIR, _MATH_DIR = _TMP / "tairdoc_icons", _TMP / "tairdoc_math"


# --------------------------------------------------------------------------- low-level xml
def _set_rfonts(rpr, name):
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for a in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        rf.attrib.pop(qn(f"w:{a}"), None)
    for a in ("ascii", "hAnsi", "eastAsia", "cs"):
        rf.set(qn(f"w:{a}"), name)


def _set_color(rpr, hexcolor=INK_HEX):
    c = rpr.find(qn("w:color"))
    if c is None:
        c = OxmlElement("w:color")
        rpr.append(c)
    for a in ("themeColor", "themeShade", "themeTint"):
        c.attrib.pop(qn(f"w:{a}"), None)
    c.set(qn("w:val"), hexcolor)


def _track(rpr, pt):
    """Letter-spacing in points (tracking for capitals)."""
    sp = OxmlElement("w:spacing")
    sp.set(qn("w:val"), str(int(round(pt * 20))))
    rpr.append(sp)


def _run(p, text, font=None, size=None, bold=None, italic=None, color=None, sup=False, track=0):
    r = p.add_run(text)
    rpr = r._element.get_or_add_rPr()
    if font:
        _set_rfonts(rpr, font)
    if size:
        r.font.size = Pt(size)
    if bold is not None:
        r.bold = bold
    if italic is not None:
        r.italic = italic
    if color:
        _set_color(rpr, color)
    if sup:
        r.font.superscript = True
    if track:
        _track(rpr, track)
    return r


def _ppr_append(p, el):
    p._p.get_or_add_pPr().append(el)


def _no_hyphen(p):
    _ppr_append(p, OxmlElement("w:suppressAutoHyphens"))


def _pborder(p, edge="left", color=CODE_BAR_HEX, sz=6, space=8):
    pPr = p._p.get_or_add_pPr()
    bdr = pPr.find(qn("w:pBdr"))
    if bdr is None:
        bdr = OxmlElement("w:pBdr")
        pPr.append(bdr)
    el = OxmlElement(f"w:{edge}")
    for k, v in (("val", "single"), ("sz", str(sz)), ("space", str(space)), ("color", color)):
        el.set(qn(f"w:{k}"), v)
    bdr.append(el)


def _cell(cell, borders=None, margins=(0, 0, 0, 0), valign=None):
    """borders: {edge: (eighths_of_pt, hex)}; unspecified edges get no line.  margins in twips."""
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("w:tcBorders", "w:tcMar", "w:vAlign"):
        old = tcPr.find(qn(tag))
        if old is not None:
            tcPr.remove(old)
    b = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        spec = (borders or {}).get(edge)
        if spec:
            for k, v in (("val", "single"), ("sz", str(spec[0])), ("space", "0"), ("color", spec[1])):
                el.set(qn(f"w:{k}"), v)
        else:
            el.set(qn("w:val"), "nil")
        b.append(el)
    tcPr.append(b)
    mar = OxmlElement("w:tcMar")
    for edge, v in zip(("top", "left", "bottom", "right"), margins):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:w"), str(int(v)))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tcPr.append(mar)
    if valign is not None:
        cell.vertical_alignment = valign


def _layout(table, widths_cm, indent_cm=0.0):
    tblPr = table._tbl.tblPr
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tblPr.append(lay)
    tw = OxmlElement("w:tblW")
    tw.set(qn("w:w"), str(int(sum(widths_cm) * 567)))
    tw.set(qn("w:type"), "dxa")
    old = tblPr.find(qn("w:tblW"))
    if old is not None:
        tblPr.remove(old)
    tblPr.append(tw)
    if indent_cm:
        ind = OxmlElement("w:tblInd")
        ind.set(qn("w:w"), str(int(indent_cm * 567)))
        ind.set(qn("w:type"), "dxa")
        tblPr.append(ind)
    for i, w in enumerate(widths_cm):
        table.columns[i].width = Cm(w)
        for c in table.columns[i].cells:
            c.width = Cm(w)


def _no_split(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:cantSplit"))


# --------------------------------------------------------------------------- figures (matplotlib)
def _mpl():
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    return matplotlib, plt


def mpl_style(theme="article"):
    """Make matplotlib match the document: same serif, CM math, black/grey ink, thin axes."""
    matplotlib, _ = _mpl()
    from matplotlib import font_manager
    want = THEMES[theme]["serif"]
    have = {f.name for f in font_manager.fontManager.ttflist}
    serif = [n for n in FALLBACKS.get(want, [want]) + ["Georgia", "DejaVu Serif"] if n in have]
    matplotlib.rcParams.update({
        "font.family": "serif", "font.serif": serif or ["DejaVu Serif"], "font.size": 9.5,
        "mathtext.fontset": THEMES[theme]["math"],
        "text.color": "#111111", "axes.edgecolor": "#111111", "axes.labelcolor": "#111111",
        "xtick.color": "#111111", "ytick.color": "#111111",
        "axes.spines.top": False, "axes.spines.right": False, "axes.linewidth": 0.7,
        "xtick.major.width": 0.7, "ytick.major.width": 0.7, "xtick.direction": "out",
        "lines.linewidth": 1.2, "legend.frameon": False, "axes.grid": False,
        "axes.prop_cycle": matplotlib.cycler(color=["#111111", "#777777", "#b0b0b0"],
                                             linestyle=["-", "--", ":"]),
        "savefig.dpi": 300, "savefig.bbox": "tight", "savefig.facecolor": "white",
    })


class Diagram:
    """Small TikZ-like diagrams: thin black boxes, serif labels, light arrowheads.

        g = Diagram(15, 3.2, theme="article")
        g.node("w", 2.2, 1.6, "Weather", "Dry / Rainy")
        g.node("e", 7.5, 1.6, "Path", "Obstacle / Clear")
        g.edge("w", "e", "influences")
        g.save("bn.png")
    Coordinates are centimetres on the final figure (width_cm × height_cm).
    """

    def __init__(self, width_cm, height_cm, theme="article"):
        mpl_style(theme)
        _, plt = _mpl()
        self.plt = plt
        self.fig = plt.figure(figsize=(width_cm / 2.54, height_cm / 2.54))
        self.ax = self.fig.add_axes([0, 0, 1, 1])
        self.ax.set_xlim(0, width_cm)
        self.ax.set_ylim(0, height_cm)
        self.ax.set_aspect("equal")
        self.ax.axis("off")
        self.nodes = {}

    def node(self, key, x, y, text, sub=None, shape="box", w=None, h=None, emphasis=False):
        from matplotlib.patches import Ellipse, FancyBboxPatch
        w = w or max(len(text) * 0.2, len(sub or "") * 0.17) + 0.8
        h = h or (1.15 if sub else 0.72)
        lw = 1.1 if emphasis else 0.75
        if shape == "circle":
            patch = Ellipse((x, y), w, h, fc="white", ec="#111111", lw=lw)
        elif shape == "plain":
            patch = FancyBboxPatch((x - w / 2, y - h / 2), w, h, boxstyle="square,pad=0", fc="none", ec="none")
        else:
            patch = FancyBboxPatch((x - w / 2, y - h / 2), w, h, boxstyle="round,pad=0,rounding_size=0.07",
                                   fc="white", ec="#111111", lw=lw)
        self.ax.add_patch(patch)
        if sub:
            self.ax.text(x, y + 0.2, text, ha="center", va="center", fontsize=10.5)
            self.ax.text(x, y - 0.24, sub, ha="center", va="center", fontsize=8.5, style="italic", color="#444444")
        else:
            self.ax.text(x, y, text, ha="center", va="center", fontsize=10.5)
        self.nodes[key] = (patch, x, y, w, h)
        return patch

    def edge(self, a, b, label=None, rad=0.0, side=1, offset=0.28, dashed=False, label_at=0.5):
        from matplotlib.patches import FancyArrowPatch
        pa, xa, ya, wa, ha = self.nodes[a]
        pb, xb, yb, wb, hb = self.nodes[b]
        arr = FancyArrowPatch((xa, ya), (xb, yb), patchA=pa, patchB=pb, shrinkA=1.5, shrinkB=1.5,
                              arrowstyle="-|>,head_length=0.38,head_width=0.16", mutation_scale=10,
                              lw=0.8, color="#111111", linestyle="--" if dashed else "-",
                              connectionstyle=f"arc3,rad={rad}")
        self.ax.add_patch(arr)
        if label:
            dx, dy = xb - xa, yb - ya
            n = max((dx * dx + dy * dy) ** 0.5, 1e-9)

            def exit_t(w, h):  # fraction of the centre-to-centre line that lies inside a w×h box
                tx = (w / 2) / abs(dx) if dx else float("inf")
                ty = (h / 2) / abs(dy) if dy else float("inf")
                return min(tx, ty)
            t0, t1 = exit_t(wa, ha), 1 - exit_t(wb, hb)
            t = t0 + (t1 - t0) * label_at          # label sits on the visible part of the arrow
            mx, my = xa + dx * t, ya + dy * t
            nx, ny = -dy / n, dx / n
            shift = side * offset - rad * n * 0.5
            self.ax.text(mx + nx * shift, my + ny * shift, label, ha="center", va="center",
                         fontsize=8.5, style="italic")
        return arr

    def text(self, x, y, s, **kw):
        kw.setdefault("fontsize", 8.5)
        kw.setdefault("style", "italic")
        return self.ax.text(x, y, s, ha=kw.pop("ha", "center"), va=kw.pop("va", "center"), **kw)

    def save(self, path):
        self.fig.savefig(path, dpi=300, bbox_inches="tight", pad_inches=0.04, facecolor="white")
        self.plt.close(self.fig)
        return path


def render_terminal(text, out_png, command=None, width_cm=15.0):
    """Program output as a quiet 'screenshot' that belongs in the document: white, hairline frame, mono."""
    _, plt = _mpl()
    from matplotlib import font_manager
    from matplotlib.patches import Rectangle
    have = {f.name for f in font_manager.fontManager.ttflist}
    mono = MONO if MONO in have else "DejaVu Sans Mono"
    lines = ([f"$ {command}"] if command else []) + text.rstrip("\n").split("\n")
    lh = 0.43  # cm per line
    h = lh * len(lines) + 0.7
    fig = plt.figure(figsize=(width_cm / 2.54, h / 2.54))
    ax = fig.add_axes([0, 0, 1, 1])
    ax.set_xlim(0, width_cm); ax.set_ylim(0, h); ax.axis("off")
    ax.add_patch(Rectangle((0.02, 0.02), width_cm - 0.04, h - 0.04, fc="#fafaf8", ec="#9a9a9a", lw=0.6))
    for i, ln in enumerate(lines):
        ax.text(0.45, h - 0.5 - i * lh, ln, family=mono, fontsize=8.4, va="center",
                color="#6a6a6a" if (command and i == 0) else "#111111")
    fig.savefig(out_png, dpi=300, facecolor="white")
    plt.close(fig)
    return out_png


def _math_png(tex, size, fontset):
    matplotlib, plt = _mpl()
    _MATH_DIR.mkdir(parents=True, exist_ok=True)
    out = _MATH_DIR / (hashlib.md5(f"{tex}|{size}|{fontset}".encode()).hexdigest()[:16] + ".png")
    if not out.exists():
        with matplotlib.rc_context({"mathtext.fontset": fontset}):
            fig = plt.figure(figsize=(0.01, 0.01))
            fig.text(0, 0, f"${tex}$", fontsize=size, color="#111111")
            fig.savefig(out, dpi=300, bbox_inches="tight", pad_inches=0.012, transparent=True)
            plt.close(fig)
    from PIL import Image
    with Image.open(out) as im:
        w, h = im.size
    return out, w / 300.0, h / 300.0


def _smart(s):
    """Typographer's quotes and dashes for prose: 'x' → ‘x’, "x" → “x”, -- → –, --- → —."""
    s = s.replace("---", "\u2014").replace("--", "\u2013")
    s = s.replace(" | ", "\u00a0|\u00a0")          # keep P(A | B) on one line
    s = re.sub(r'(^|[\s(\[\u2014\u2013/])"', "\\1\u201c", s)
    s = s.replace('"', "\u201d")
    s = re.sub(r"(^|[\s(\[\u2014\u2013/])'", "\\1\u2018", s)
    return s.replace("'", "\u2019")


def _icon_png(kind):
    _, plt = _mpl()
    from matplotlib.patches import Circle
    glyph = {"info": "i", "note": "i", "tip": "\u2713", "answer": "\u2713", "warning": "!"}.get(kind, "i")
    _ICON_DIR.mkdir(parents=True, exist_ok=True)
    out = _ICON_DIR / f"{kind}.png"
    if not out.exists():
        fig = plt.figure(figsize=(0.4, 0.4))
        ax = fig.add_axes([0, 0, 1, 1]); ax.set_xlim(0, 1); ax.set_ylim(0, 1); ax.axis("off")
        ax.add_patch(Circle((0.5, 0.5), 0.47, color="#111111"))
        ax.text(0.5, 0.48, glyph, ha="center", va="center", color="white", fontsize=17, weight="bold",
                family="DejaVu Serif" if glyph == "i" else "DejaVu Sans")
        fig.savefig(out, dpi=300, transparent=True)
        plt.close(fig)
    return out


# --------------------------------------------------------------------------- inline markup
_NOTE = re.compile(r"\^\[((?:[^\[\]]|\[[^\]]*\])+)\]")
_TOK = re.compile(r"(\x00\d+\x00|\*\*[^*]+\*\*|`[^`]+`|\$[^$]+\$|\*[^*\s][^*]*\*)")


# --------------------------------------------------------------------------- the document
class Doc:
    def __init__(self, title, author=None, email=None, affiliation=None, date=None, subtitle=None,
                 theme="article", running=None, keywords=None):
        if theme not in THEMES:
            raise ValueError(f"theme must be one of {list(THEMES)}")
        self.theme = theme
        self.t = THEMES[theme]
        self.serif = pick_font(self.t["serif"])
        self.prose_font = f"{self.serif}:{self.t['prose']}" if self.t["prose"] else self.serif
        self.fig_font = f"{self.serif}:{self.t['figs']}"
        self.smcp_font = f"{self.serif}:smcp" + (f"&{self.t['prose']}" if self.t["prose"] else "")
        L, R, T, B = self.t["margins"]
        self.text_w = 21.0 - L - R
        self.running = running or title
        self.d = Document()
        self._n = dict(sec=0, sub=0, fig=0, tab=0, eq=0, lst=0, note=0)
        self._sec_numbered = True
        self._indent_next = False
        self._headings = []      # (level, title) for contents()
        self._toc_anchor = None
        self._page()
        self._styles()
        self._title_block(title, subtitle, author, email, affiliation, date, keywords)

    # ................................................................ page + styles
    def _page(self):
        L, R, T, B = self.t["margins"]
        s = self.d.sections[0]
        s.page_width, s.page_height = Cm(21.0), Cm(29.7)
        s.left_margin, s.right_margin, s.top_margin, s.bottom_margin = Cm(L), Cm(R), Cm(T), Cm(B)
        s.footer_distance = Cm(1.25)
        fp = s.footer.paragraphs[0]
        mode = self.t["footer"]
        size = self.t["body"] - 1
        if mode == "running":
            fp.paragraph_format.tab_stops.add_tab_stop(Cm(21.0 - L - R), WD_TAB_ALIGNMENT.RIGHT)
            _pborder(fp, "top", RULE_HEX, 4, 4)
            _run(fp, self.running, self.smcp_font, size)
            _run(fp, "\t", size=size)
        elif mode == "right":
            fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        else:
            fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fld = OxmlElement("w:fldSimple")
        fld.set(qn("w:instr"), "PAGE")
        r = OxmlElement("w:r")
        rpr = OxmlElement("w:rPr")
        sz = OxmlElement("w:sz"); sz.set(qn("w:val"), str(int(size * 2)))
        rpr.append(sz)
        _set_rfonts(rpr, self.prose_font)
        t = OxmlElement("w:t"); t.text = "1"
        r.append(rpr); r.append(t); fld.append(r)
        fp._p.append(fld)
        # Justified text needs hyphenation to avoid rivers; LaTeX does this by default.
        st = self.d.settings.element
        ah = OxmlElement("w:autoHyphenation")
        zoom = st.find(qn("w:zoom"))
        (zoom.addnext(ah) if zoom is not None else st.insert(0, ah))

    def _styles(self):
        styles = self.d.styles
        dd = styles.element.find(qn("w:docDefaults"))
        if dd is not None:
            rpr = dd.find(qn("w:rPrDefault")).find(qn("w:rPr"))
            _set_rfonts(rpr, self.prose_font)
            lang = rpr.find(qn("w:lang"))
            if lang is None:
                lang = OxmlElement("w:lang"); rpr.append(lang)
            lang.set(qn("w:val"), "en-GB")
        n = styles["Normal"]
        _set_rfonts(n.element.get_or_add_rPr(), self.prose_font)
        n.font.size = Pt(self.t["body"])
        n.font.color.rgb = INK
        pf = n.paragraph_format
        pf.space_before = pf.space_after = Pt(0)
        pf.line_spacing = self.t["leading"]
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf.widow_control = True
        for lvl in (1, 2, 3):
            spec = self.t[f"h{lvl}"]
            h = styles[f"Heading {lvl}"]
            rpr = h.element.get_or_add_rPr()
            _set_rfonts(rpr, self.serif)
            _set_color(rpr)
            h.font.size, h.font.bold, h.font.italic = Pt(spec["size"]), spec["bold"], spec["italic"]
            if spec["track"]:
                _track(rpr, spec["track"])
            hp = h.paragraph_format
            hp.space_before, hp.space_after = Pt(spec["before"]), Pt(spec["after"])
            hp.keep_with_next, hp.alignment, hp.line_spacing = True, WD_ALIGN_PARAGRAPH.LEFT, 1.0
            _no_hyphen_style(h)

    # ................................................................ title block
    def _line(self, text, size, font=None, italic=False, before=0, after=0, align=None, track=0, bold=False):
        p = self.d.add_paragraph()
        p.alignment = align if align is not None else (
            WD_ALIGN_PARAGRAPH.CENTER if self.t["title"]["align"] == "center" else WD_ALIGN_PARAGRAPH.LEFT)
        p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(before), Pt(after)
        p.paragraph_format.line_spacing = 1.05
        _no_hyphen(p)
        _run(p, _smart(text) if font != MONO else text, font or self.prose_font, size,
             italic=italic, track=track, bold=bold)
        return p

    def _rule(self, before=0, after=0, sz=6):
        p = self.d.add_paragraph()
        p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(before), Pt(after)
        p.paragraph_format.line_spacing = Pt(1)
        _pborder(p, "bottom", RULE_HEX, sz, 0)
        return p

    def _title_block(self, title, subtitle, author, email, affiliation, date, keywords):
        th = self.t["title"]
        if self.theme == "journal":
            self._rule(after=14, sz=6)
            self._line(title, th["size"], after=10)
            if subtitle:
                self._line(subtitle, 12.5, italic=True, after=8)
            if author:
                self._line(author, 11.5, self.smcp_font, after=5)
            for extra in (affiliation, email and f"Email: {email}", date):
                if extra:
                    self._line(extra, 9.3, italic=True, after=1)
            self._keywords = keywords
            self._pending_rule = True        # closed after abstract(), or before first section
        elif self.theme == "essay":
            self._line(title, th["size"], italic=True, before=4, after=4)
            if subtitle:
                self._line(subtitle, 13, italic=True, after=2)
            for extra in (author, date):
                if extra:
                    self._line(extra, 13, italic=True, after=1)
            if affiliation or email:
                self._line(" \u00b7 ".join(x for x in (affiliation, email) if x), 9.5, italic=False, after=0,
                           font=self.smcp_font, track=0.3)
            self._spacer(16)
        else:
            self._line(title, th["size"], before=4, after=10)
            if subtitle:
                self._line(subtitle, 12.5, italic=True, after=8)
            if author:
                self._line(author, 12, after=1)
            if email:
                self._line(email, 10, MONO, after=1)
            tail = " \u2014 ".join(x for x in (affiliation, date) if x)
            if tail:
                self._line(tail, 12, before=9)
            self._spacer(18)
        self._indent_next = False

    def _close_title(self):
        if getattr(self, "_pending_rule", False):
            if self._keywords:
                p = self._line("", 9.3, after=0)
                _run(p, "Keywords: ", self.prose_font, 9.3, italic=True)
                _run(p, "; ".join(self._keywords), self.prose_font, 9.3, italic=True)
                p.paragraph_format.space_before = Pt(8)
            self._rule(before=10, after=6, sz=6)
            self._pending_rule = False

    # ................................................................ helpers
    def _spacer(self, pt):
        p = self.d.add_paragraph()
        p.paragraph_format.line_spacing = Pt(max(pt, 1))
        return p

    def _rich(self, p, text, size=None, font=None, italic=None, notes=None):
        font = font or self.prose_font
        for part in _TOK.split(text):
            if not part:
                continue
            if part.startswith("\x00"):
                _run(p, part.strip("\x00"), font, (size or self.t["body"]) - 0.5, sup=True)
            elif part.startswith("**") and part.endswith("**") and len(part) > 4:
                _run(p, _smart(part[2:-2]), font, size, bold=True, italic=italic)
            elif part.startswith("`") and part.endswith("`") and len(part) > 2:
                _run(p, part[1:-1], MONO, (size or self.t["body"]) - 1.2)
            elif part.startswith("$") and part.endswith("$") and len(part) > 2:
                _run(p, part[1:-1], self.fig_font, size, italic=True)
            elif part.startswith("*") and part.endswith("*") and len(part) > 2:
                _run(p, _smart(part[1:-1]), font, size, italic=not italic)
            else:
                _run(p, _smart(part), font, size, italic=italic)

    def _take_notes(self, text):
        notes = []

        def sub(m):
            self._n["note"] += 1
            notes.append((self._n["note"], m.group(1).strip()))
            return f"\x00{self._n['note']}\x00"
        return _NOTE.sub(sub, text), notes

    def _note_paragraphs(self, container, notes, size=None, first_para=None):
        size = size or (8.6 if self.theme == "essay" else 8.8)
        for i, (n, txt) in enumerate(notes):
            p = first_para if (i == 0 and first_para is not None) else container.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.line_spacing = 1.12
            p.paragraph_format.space_after = Pt(5)
            _no_hyphen(p)
            _run(p, str(n), self.prose_font, size - 1.5, sup=True)
            _run(p, "\u2009", self.prose_font, size)
            self._rich(p, txt, size=size)

    def _with_margin(self, build_main, notes=None, valign_notes=WD_CELL_VERTICAL_ALIGNMENT.TOP):
        """Two-cell borderless row: main column + right margin column (essay theme)."""
        t = self.d.add_table(rows=1, cols=2)
        t.alignment = WD_TABLE_ALIGNMENT.LEFT
        gap, col = self.t["note_gap"], self.t["note_col"]
        _layout(t, [self.text_w, gap + col])
        main, side = t.cell(0, 0), t.cell(0, 1)
        _cell(main, margins=(0, 0, 0, 0), valign=WD_CELL_VERTICAL_ALIGNMENT.TOP)
        _cell(side, margins=(20, int(gap * 567), 0, 0), valign=valign_notes)
        build_main(main)
        return t, side

    # ................................................................ structure
    def section(self, title, numbered=True):
        self._close_title()
        spec = self.t["h1"]
        self._sec_numbered = numbered
        if numbered:
            self._n["sec"] += 1
            self._n["sub"] = 0
        shown = spec["fmt"].format(n=self._n["sec"], t=title) if numbered else title
        if spec["upper"]:
            shown = shown.upper()
        self.d.add_heading(shown, level=1)
        self._headings.append((1, title if not numbered else f"{self._n['sec']}\u2003{title}"))
        self._indent_next = False

    def subsection(self, title, numbered=None):
        spec = self.t["h2"]
        numbered = self._sec_numbered if numbered is None else numbered
        if numbered and self._n["sec"]:
            self._n["sub"] += 1
            num = f"{self._n['sec']}.{self._n['sub']}"
            shown = spec["fmt"].format(n=num, t=title)
            toc = f"{num}\u2003{title}"
        else:
            shown = toc = title
        self.d.add_heading(shown, level=2)
        self._headings.append((2, toc))
        self._indent_next = False

    def subsubsection(self, title):
        self.d.add_heading(title, level=3)
        self._indent_next = False

    def abstract(self, text, label=None):
        """Article/journal: indented block under the title.  Essay: 'Summary' run-in block."""
        if self.theme == "essay":
            p = self.d.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.55)
            p.paragraph_format.right_indent = Cm(0.8)
            p.paragraph_format.space_after = Pt(14)
            p.paragraph_format.line_spacing = 1.22
            _run(p, (label or "Summary") + " ", self.prose_font, 9.6, bold=True)
            self._rich(p, text, size=9.6)
        else:
            if self.theme == "article":
                hp = self._line(label or "Abstract", 10, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=3)
            p = self.d.add_paragraph()
            ind = 1.0 if self.theme == "journal" else 1.3
            p.paragraph_format.left_indent = p.paragraph_format.right_indent = Cm(ind)
            p.paragraph_format.space_after = Pt(12 if self.theme == "article" else 0)
            p.paragraph_format.line_spacing = 1.12
            if self.theme == "journal":
                p.paragraph_format.space_before = Pt(10)
            self._rich(p, text, size=9.8)
            self._close_title()
        self._indent_next = False

    # ................................................................ text
    def para(self, text, lead=None):
        """Body paragraph.  Markup: **bold**, *italic*, `mono`, $var$, ^[side/foot note].
        `lead` = bold run-in label, e.g. 'Question 1.' (LaTeX \\paragraph)."""
        self._close_title()
        text, notes = self._take_notes(text)
        indent = self._indent_next

        def fill(p):
            if indent:
                p.paragraph_format.first_line_indent = Cm(self.t["indent"])
            if lead:
                _run(p, lead.rstrip() + " ", self.prose_font, bold=True)
            self._rich(p, text)

        if notes and self.t["note_col"]:
            _, side = self._with_margin(lambda c: fill(c.paragraphs[0]))
            self._note_paragraphs(side, notes, first_para=side.paragraphs[0])
        else:
            p = self.d.add_paragraph()
            fill(p)
            if notes:  # article/journal: notes as a small block right under the paragraph
                p.paragraph_format.space_after = Pt(3)
                box = self.d.add_paragraph()
                box.paragraph_format.left_indent = Cm(self.t["indent"])
                box.paragraph_format.space_after = Pt(4)
                self._note_paragraphs(self.d, notes, first_para=box)
        self._indent_next = True

    def bullets(self, items, numbered=False):
        """items: strings, or (lead, text) tuples for a bold run-in label."""
        self._close_title()
        L, hang = 0.95, 0.45
        for i, it in enumerate(items, 1):
            p = self.d.add_paragraph()
            pf = p.paragraph_format
            pf.left_indent, pf.first_line_indent = Cm(L), Cm(-hang)
            pf.tab_stops.add_tab_stop(Cm(L))
            pf.space_after = Pt(2.5)
            mark = f"{i}." if numbered else self.t["bullet"]
            _run(p, mark + "\t", self.fig_font if numbered else self.prose_font)
            if isinstance(it, (tuple, list)):
                _run(p, it[0].rstrip() + " ", self.prose_font, bold=True)
                it = it[1]
            text, notes = self._take_notes(it)
            self._rich(p, text)
        self.d.paragraphs[-1].paragraph_format.space_after = Pt(7)
        self._indent_next = False

    def quote(self, text, source=None):
        self._close_title()
        p = self.d.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent = pf.right_indent = Cm(1.0)
        pf.space_before, pf.space_after = Pt(8), Pt(8)
        self._rich(p, text, size=self.t["body"] - 0.5)
        if source:
            _run(p, f" \u2014\u2009{source}", self.prose_font, self.t["body"] - 0.5)
        self._indent_next = False

    def equation(self, tex, numbered=True, size=None):
        """Display maths in Computer Modern, numbered flush right like LaTeX.  Returns the number."""
        self._close_title()
        size = size or self.t["body"] + 1.5
        tex = re.sub(r"(?<![a-zA-Z])\\frac(?=\s*\{)", r"\\dfrac", tex)   # display-size fractions
        png, w_in, h_in = _math_png(tex, size, self.t["math"])
        room = (self.text_w - (1.7 if numbered else 0.3)) / 2.54     # leave space for "(n)"
        if w_in > room:                                               # too wide: shrink to fit
            h_in *= room / w_in
            w_in = room
        p = self.d.add_paragraph()
        pf = p.paragraph_format
        pf.space_before, pf.space_after, pf.line_spacing = Pt(7), Pt(7), 1.0
        _no_hyphen(p)
        num = None
        pic = lambda: p.add_run().add_picture(str(png), width=Emu(int(w_in * 914400)))
        if numbered:
            self._n["eq"] += 1
            num = self._n["eq"]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            pf.tab_stops.add_tab_stop(Cm(self.text_w / 2), WD_TAB_ALIGNMENT.CENTER)
            pf.tab_stops.add_tab_stop(Cm(self.text_w), WD_TAB_ALIGNMENT.RIGHT)
            p.add_run("\t"); pic()
            r = _run(p, f"\t({num})", self.fig_font)
            lift = (h_in * 72 - self.t["body"] * 0.7) / 2      # centre the number on the formula
            if lift > 0.5:
                pos = OxmlElement("w:position")
                pos.set(qn("w:val"), str(int(lift * 2)))
                r._element.get_or_add_rPr().append(pos)
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pic()
        self._indent_next = False
        return num

    def callout(self, text, kind="info", label=None):
        """Aside with icon + left rule.  kind: info | note | tip | answer | warning.  text: str or list."""
        self._close_title()
        label = label or {"info": "Info", "note": "Note", "tip": "Tip", "answer": "Answer",
                          "warning": "Warning"}.get(kind, kind.title())
        self._spacer(5)
        t = self.d.add_table(rows=1, cols=2)
        t.alignment = WD_TABLE_ALIGNMENT.LEFT
        _layout(t, [0.62, self.text_w - 0.62])
        icon, body = t.cell(0, 0), t.cell(0, 1)
        _cell(icon, margins=(10, 0, 0, 45), valign=WD_CELL_VERTICAL_ALIGNMENT.TOP)
        _cell(body, borders={"left": (9, RULE_HEX)}, margins=(15, 210, 15, 40))
        _no_split(t.rows[0])
        ip = icon.paragraphs[0]
        ip.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        ip.add_run().add_picture(str(_icon_png(kind)), width=Cm(0.33))
        size = self.t["body"] - 0.5
        for i, chunk in enumerate(text if isinstance(text, (list, tuple)) else [text]):
            bp = body.paragraphs[0] if i == 0 else body.add_paragraph()
            bp.paragraph_format.space_after = Pt(3)
            if i == 0:
                _run(bp, f"{label}: ", self.prose_font, size, bold=True)
            chunk, _ = self._take_notes(chunk)
            self._rich(bp, chunk, size=size)
        self._spacer(7)
        self._indent_next = False

    def code(self, text, caption=None, size=8.5):
        self._close_title()
        if caption:
            self._n["lst"] += 1
            c = self.d.add_paragraph()
            c.paragraph_format.space_before = Pt(8)
            c.paragraph_format.keep_with_next = True
            self._caption_runs(c, "Listing", self._n["lst"], caption)
        lines = text.expandtabs(4).rstrip("\n").split("\n")
        longest = max(len(ln) for ln in lines)
        fit = (self.text_w - 0.8) / (longest * 0.55 * 0.03528) if longest else size  # Consolas ≈ 0.55 em
        size = max(6.8, min(size, fit))       # shrink long listings instead of wrapping lines
        for i, ln in enumerate(lines):          # one paragraph per line; identical borders merge into one bar
            p = self.d.add_paragraph()
            pf = p.paragraph_format
            pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
            pf.left_indent = Cm(0.35)
            pf.line_spacing = 1.0
            pf.space_before = Pt((4 if caption else 9) if i == 0 else 0)
            pf.space_after = Pt(9 if i == len(lines) - 1 else 0)
            pf.keep_with_next = i < len(lines) - 1 and len(lines) <= 25
            _no_hyphen(p)
            _pborder(p, "left", CODE_BAR_HEX, 6, 8)
            _run(p, ln or " ", MONO, size)
        self._indent_next = False

    # ................................................................ floats
    def _caption_runs(self, p, kind, n, text, size=None):
        cap = self.t["caption"]
        size = size or cap["size"]
        label = cap["label"].format(K=kind, n=n)
        if cap["smcp"]:
            _run(p, label + " ", self.smcp_font, size)
        else:
            _run(p, label + " ", self.prose_font, size, italic=cap["italic"])
        self._rich(p, text, size=size, italic=cap["italic"] or None)

    def figure(self, path, caption=None, width_cm=None, wide=False):
        """Centred figure with caption below.  Essay theme puts the caption in the margin
        (unless wide=True, which spans text + margin)."""
        self._close_title()
        from PIL import Image
        with Image.open(path) as im:
            w_px = im.size[0]
            dpi = (im.info.get("dpi") or (300, 300))[0] or 300
        natural = w_px / dpi * 2.54
        if self.t["note_col"] and not width_cm and natural > self.text_w + 1.5:
            wide = True                     # essay: let broad diagrams run into the margin
        full = self.text_w + (self.t["note_gap"] + self.t["note_col"] if wide else 0)
        width_cm = min(width_cm or natural, full)
        if caption:
            self._n["fig"] += 1
        margin_caption = bool(caption) and self.t["note_col"] and not wide
        if margin_caption:
            def build(c):
                p = c.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(str(path), width=Cm(min(width_cm, self.text_w)))
            self._spacer(6)
            t, side = self._with_margin(build, valign_notes=WD_CELL_VERTICAL_ALIGNMENT.BOTTOM)
            _no_split(t.rows[0])
            sp = side.paragraphs[0]
            sp.alignment = WD_ALIGN_PARAGRAPH.LEFT
            sp.paragraph_format.line_spacing = 1.12
            _no_hyphen(sp)
            self._caption_runs(sp, "Figure", self._n["fig"], caption, size=8.6)
            self._spacer(8)
        else:
            p = self.d.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if not wide else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.keep_with_next = bool(caption)
            p.paragraph_format.line_spacing = 1.0
            overhang = Cm(-(self.t["note_gap"] + self.t["note_col"])) if wide and self.t["note_col"] else None
            if overhang is not None:
                p.paragraph_format.right_indent = overhang
            p.add_run().add_picture(str(path), width=Cm(width_cm))
            if caption:
                c = self.d.add_paragraph()
                c.alignment = WD_ALIGN_PARAGRAPH.CENTER
                c.paragraph_format.space_before, c.paragraph_format.space_after = Pt(5), Pt(12)
                _no_hyphen(c)
                if wide:
                    c.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    if overhang is not None:
                        c.paragraph_format.right_indent = overhang
                self._caption_runs(c, "Figure", self._n["fig"], caption)
            else:
                p.paragraph_format.space_after = Pt(12)
        self._indent_next = False
        return self._n["fig"]

    def table(self, rows, caption=None, header=True, align=None, widths=None, size=None, full=False):
        """Booktabs table: heavy top/bottom rule, light rule under the header, nothing else.

        rows   list of lists; first row is the header when header=True.
        align  e.g. "lcc"; default: numeric columns centred, text left.
        widths list of cm; default sized to content (natural width, like LaTeX).  full=True → text width.
        """
        self._close_title()
        size = size or (self.t["body"] - 1)
        ncol = len(rows[0])
        if caption:
            self._n["tab"] += 1
            c = self.d.add_paragraph()
            c.alignment = WD_ALIGN_PARAGRAPH.CENTER
            c.paragraph_format.space_before, c.paragraph_format.space_after = Pt(9), Pt(4)
            c.paragraph_format.keep_with_next = True
            _no_hyphen(c)
            self._caption_runs(c, "Table", self._n["tab"], caption)
        else:
            self._spacer(6)
        num = re.compile(r"^[\s\-+\u2212\u2248~(]*[\d.,%/\u00d7]+[\s%)]*$")
        if align is None:
            align = ""
            for j in range(ncol):
                body = [str(r[j]) for r in rows[1 if header else 0:] if str(r[j]).strip()]
                align += "c" if body and all(num.match(v) for v in body) else "l"
        if widths is None:
            widths = self._col_widths(rows, header, size, full)
        t = self.d.add_table(rows=len(rows), cols=ncol)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        _layout(t, widths)
        heavy, light, last = (11, RULE_HEX), (4, RULE_HEX), len(rows) - 1
        amap = {"l": WD_ALIGN_PARAGRAPH.LEFT, "c": WD_ALIGN_PARAGRAPH.CENTER, "r": WD_ALIGN_PARAGRAPH.RIGHT}
        head = self.t["table_head"]
        for i, row in enumerate(rows):
            _no_split(t.rows[i])
            if header and i == 0:
                t.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
            for j, val in enumerate(row):
                cell = t.cell(i, j)
                b = {}
                if i == 0:
                    b["top"] = heavy
                if header and i == 0:
                    b["bottom"] = light
                if i == last:
                    b["bottom"] = heavy
                top = 70 if i == 0 or (header and i == 1) else 25
                bot = 70 if i == last or (header and i == 0) else 25
                _cell(cell, borders=b, margins=(top, 100, bot, 100), valign=WD_CELL_VERTICAL_ALIGNMENT.TOP)
                p = cell.paragraphs[0]
                p.alignment = amap.get(align[j], WD_ALIGN_PARAGRAPH.LEFT)
                p.paragraph_format.line_spacing = 1.0
                if i < last and len(rows) <= 18:          # short tables never split across pages
                    p.paragraph_format.keep_with_next = True
                _no_hyphen(p)
                if header and i == 0:
                    if head == "smcp":
                        _run(p, str(val), self.smcp_font, size)
                    else:
                        _run(p, str(val), self.serif, size, bold=head == "bold", italic=head == "italic")
                else:
                    self._rich(p, str(val), size=size, font=self.fig_font)
        self._spacer(9)
        self._indent_next = False
        return self._n["tab"]

    def _col_widths(self, rows, header, size, full):
        """Natural widths; if they don't fit, short columns keep theirs and only long ones wrap."""
        per_char = 0.185 * size / 10                      # cm per character at this size
        nat = []
        for j in range(len(rows[0])):
            lens = [len(str(r[j])) * (1.12 if header and i == 0 else 1.0) for i, r in enumerate(rows)]
            nat.append(max(max(lens) * per_char + 0.5, 1.2))
        if sum(nat) <= self.text_w:
            return [w * self.text_w / sum(nat) for w in nat] if full else nat
        widths, left, free = [None] * len(nat), self.text_w, list(range(len(nat)))
        for j in sorted(free, key=lambda k: nat[k]):
            share = left / len(free)
            if nat[j] <= share:
                widths[j] = nat[j]
                left -= nat[j]
                free.remove(j)
            else:
                break
        for j in free:
            widths[j] = left / len(free)
        return widths

    def pagebreak(self):
        self.d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        self._indent_next = False

    def contents(self, title="Contents", depth=2):
        """Book-style contents page (filled in at save time with real page numbers)."""
        self._close_title()
        self._toc_anchor = self.d.add_paragraph()
        self._toc_depth = depth
        self._toc_title = title

    # ................................................................ contents builder
    def _build_toc(self, pages=None):
        anchor = self._toc_anchor
        entries = [(l, t) for l, t in self._headings if l <= self._toc_depth]
        made = []
        h = self.d.add_paragraph()
        h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        h.paragraph_format.space_before, h.paragraph_format.space_after = Pt(30), Pt(26)
        _run(h, self._toc_title, self.smcp_font, 13, track=0.6)
        made.append(h._p)
        tw = min(self.text_w, 12.5)
        for i, (lvl, txt) in enumerate(entries):
            p = self.d.add_paragraph()
            pf = p.paragraph_format
            pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
            pf.left_indent = Cm((self.text_w - tw) / 2 + (0.7 if lvl == 2 else 0))
            pf.space_before = Pt(7 if lvl == 1 and i else 0)
            pf.space_after = Pt(2)
            pf.tab_stops.add_tab_stop(Cm((self.text_w - tw) / 2 + tw), WD_TAB_ALIGNMENT.RIGHT)
            _no_hyphen(p)
            num = str(pages[i]) if pages and pages[i] else "00"
            _run(p, txt, self.prose_font, self.t["body"] - (0 if lvl == 1 else 0.5), italic=lvl == 2)
            _run(p, "\t" + num, f"{self.serif}:onum", self.t["body"] - 0.5)
            made.append(p._p)
        br = self.d.add_paragraph()
        br.add_run().add_break(WD_BREAK.PAGE)
        made.append(br._p)
        for el in made:
            anchor._p.addprevious(el)
        return entries

    # ................................................................ output
    def save(self, path, pdf=True):
        path = Path(path).resolve()
        self._close_title()
        if self._toc_anchor is None:
            self.d.save(path)
            return (path, to_pdf(path)) if pdf else (path, None)
        # two passes: lay out with placeholder numbers, read real pages from the PDF, refill
        body = self.d.element.body
        anchor_idx = list(body).index(self._toc_anchor._p)
        base = copy.deepcopy(body)
        entries = self._build_toc()
        self.d.save(path)
        pages = _find_heading_pages(to_pdf(path), [t for _, t in entries], self._toc_title)
        for ch in list(body):
            body.remove(ch)
        for ch in list(base):
            body.append(ch)
        self._toc_anchor = _Anchor(list(body)[anchor_idx])
        self._build_toc(pages)
        self.d.save(path)
        return path, (to_pdf(path) if pdf else None)


class _Anchor:
    def __init__(self, p_element):
        self._p = p_element


def _no_hyphen_style(style):
    ppr = style.element.get_or_add_pPr()
    ppr.append(OxmlElement("w:suppressAutoHyphens"))


def _find_heading_pages(pdf, titles, toc_title):
    import fitz
    doc = fitz.open(str(pdf))
    texts = [" ".join(pg.get_text().split()).lower() for pg in doc]
    start = 0
    for i, tx in enumerate(texts):
        if toc_title.lower() in tx:
            start = i + 1
            break
    pages, cur = [], start
    for t in titles:
        key = " ".join(t.replace("\u2003", " ").split()).lower()
        key2 = " ".join(key.split(" ")[1:]) if key[:1].isdigit() else key
        found = None
        for i in range(cur, len(texts)):
            if key in texts[i] or (key2 and key2 in texts[i]):
                found = i
                break
        if found is not None:
            cur = found
        pages.append((found + 1) if found is not None else None)
    return pages


# --------------------------------------------------------------------------- pdf + preview
def _soffice():
    for c in (shutil.which("soffice"), shutil.which("libreoffice"),
              r"C:\Program Files\LibreOffice\program\soffice.exe",
              r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
              "/Applications/LibreOffice.app/Contents/MacOS/soffice", "/usr/bin/soffice"):
        if c and Path(c).exists():
            return c
    raise RuntimeError("LibreOffice (soffice) not found: install it, or deliver the .docx")


def to_pdf(docx_path):
    docx_path = Path(docx_path).resolve()
    profile = _TMP / "tairdoc_lo_profile"
    subprocess.run([_soffice(), f"-env:UserInstallation={profile.as_uri()}", "--headless",
                    "--convert-to", "pdf", "--outdir", str(docx_path.parent), str(docx_path)],
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=240)
    pdf = docx_path.with_suffix(".pdf")
    if not pdf.exists():
        raise RuntimeError("PDF conversion failed")
    return pdf


def preview(pdf_path, out_dir=None, dpi=72, pages=None):
    """Render pages to PNG so the result can be inspected before it is delivered."""
    import fitz
    pdf_path = Path(pdf_path)
    out = Path(out_dir or pdf_path.parent / "_preview")
    out.mkdir(exist_ok=True)
    doc = fitz.open(str(pdf_path))
    paths = []
    for i in (pages or range(len(doc))):
        o = out / f"{pdf_path.stem}_p{i + 1}.png"
        doc[i].get_pixmap(dpi=dpi).save(str(o))
        paths.append(o)
    return paths
